"""Broker XLS → SQLite ingest.

Reads a broker Account-Activity export via openpyxl, computes a stable
`exec_id` per row, and bulk-inserts via `sqlite_db.insert_executions`
using INSERT OR IGNORE for idempotency.

`refresh_from_paths` is the workflow entry point: ingest → identify
affected (account, symbol) tuples → re-walk pairing per pair →
`replace_round_trips`. Used by `cli.py cmd_refresh_from_xls`.
"""

from __future__ import annotations

import glob
import os
from datetime import datetime
from pathlib import Path

import openpyxl

from .pairing import pair_executions
from .sqlite_db import (
    ExecutionRow,
    affected_symbols_for_file,
    apply_migrations,
    compute_exec_id,
    connect,
    insert_executions,
    read_executions_for_symbol,
    replace_round_trips,
)


def _f(x) -> float:
    """Coerce optional numeric XLS cell to float (None → 0.0)."""
    if x is None:
        return 0.0
    return float(x)


def parse_broker_excel_rows(path: Path, account: str) -> list[ExecutionRow]:
    """Parse a broker XLS into raw signed-quantity ExecutionRows.

    Broker exports use the .xls extension but are actually .xlsx — openpyxl
    requires .xlsx, so we copy/rename to a temp .xlsx if needed.
    """
    path = Path(path)
    load_path = path
    tmp_path: Path | None = None
    if path.suffix.lower() != ".xlsx":
        import shutil
        import tempfile

        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp.close()
        tmp_path = Path(tmp.name)
        shutil.copy(path, tmp_path)
        load_path = tmp_path

    try:
        wb = openpyxl.load_workbook(load_path, data_only=True)
        sheet = wb[wb.sheetnames[0]]
        source_file = path.name
        rows: list[ExecutionRow] = []

        for row in sheet.iter_rows(min_row=2, values_only=True):
            (
                date,
                qty,
                symbol,
                desc,
                price,
                transfee,
                ecn_taker,
                ecn_maker,
                orf,
                taf,
                sec,
                cat,
                amount,
                type_,
                commission,
            ) = row

            if (
                qty is None
                or symbol is None
                or not str(symbol).strip()
                or price is None
                or price == 0
            ):
                continue

            if isinstance(date, datetime):
                dt = date
            else:
                dt = datetime.strptime(str(date), "%Y-%m-%d")

            trade_datetime = dt.isoformat(timespec="seconds")
            trade_date = dt.strftime("%Y-%m-%d")
            signed_qty = int(qty)
            price_f = float(price)
            amount_f = float(amount) if amount is not None else 0.0
            symbol_s = str(symbol).strip()

            exec_id = compute_exec_id(
                account=account,
                trade_datetime=trade_datetime,
                symbol=symbol_s,
                quantity=signed_qty,
                price=price_f,
                amount=amount_f,
            )

            rows.append(
                ExecutionRow(
                    exec_id=exec_id,
                    account=account,
                    trade_date=trade_date,
                    trade_datetime=trade_datetime,
                    quantity=signed_qty,
                    symbol=symbol_s,
                    description=str(desc).strip() if desc else None,
                    price=price_f,
                    trans_fee=_f(transfee),
                    ecn_taker=_f(ecn_taker),
                    ecn_maker=_f(ecn_maker),
                    orf_fee=_f(orf),
                    taf_fee=_f(taf),
                    sec_fee=_f(sec),
                    cat_fee=_f(cat),
                    amount=amount_f,
                    type=str(type_) if type_ else None,
                    commissions=_f(commission),
                    source_file=source_file,
                )
            )

        wb.close()
        return rows
    finally:
        if tmp_path is not None and tmp_path.exists():
            tmp_path.unlink()


def expand_paths(patterns: list[str]) -> list[Path]:
    """Expand a mix of files and globs into a unique ordered file list."""
    out: list[Path] = []
    seen: set[Path] = set()
    for p in patterns:
        expanded = os.path.expanduser(p)
        if any(c in expanded for c in "*?[]"):
            matches = sorted(glob.glob(expanded))
        else:
            matches = [expanded]
        for m in matches:
            mp = Path(m).resolve()
            if mp in seen:
                continue
            if mp.is_file():
                seen.add(mp)
                out.append(mp)
    return out


def refresh_from_paths(
    paths: list[str],
    account: str,
    db_path: str | Path | None = None,
    dry_run: bool = False,
    verbose: bool = False,
) -> dict:
    """End-to-end: parse XLS files → ingest → repair affected RTs.

    Returns a stats dict for the CLI to report.
    """
    files = expand_paths(paths)
    if not files:
        return {
            "files": [],
            "executions_parsed": 0,
            "executions_new": 0,
            "round_trips_computed": 0,
            "dry_run": dry_run,
        }

    all_rows_by_file: list[tuple[Path, list[ExecutionRow]]] = []
    total_parsed = 0
    for f in files:
        rows = parse_broker_excel_rows(f, account=account)
        total_parsed += len(rows)
        all_rows_by_file.append((f, rows))
        if verbose:
            print(f"  {f.name}: {len(rows)} rows parsed")

    if dry_run:
        return {
            "files": [str(f) for f, _ in all_rows_by_file],
            "executions_parsed": total_parsed,
            "executions_new": 0,
            "round_trips_computed": 0,
            "dry_run": True,
        }

    conn = connect(db_path)
    try:
        apply_migrations(conn)
        total_new = 0
        affected: set[tuple[str, str]] = set()
        for f, rows in all_rows_by_file:
            attempted, new_count = insert_executions(conn, rows)
            total_new += new_count
            for pair in affected_symbols_for_file(conn, f.name):
                affected.add(pair)
            if verbose:
                print(
                    f"  {f.name}: {attempted} attempted, {new_count} newly inserted"
                )

        rt_count = 0
        for acct, sym in sorted(affected):
            execs = read_executions_for_symbol(conn, acct, sym)
            results = pair_executions(execs)
            rt_count += replace_round_trips(conn, acct, sym, results)

        return {
            "files": [str(f) for f, _ in all_rows_by_file],
            "executions_parsed": total_parsed,
            "executions_new": total_new,
            "round_trips_computed": rt_count,
            "affected_pairs": len(affected),
            "dry_run": False,
        }
    finally:
        conn.close()
